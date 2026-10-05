"""Step 3 — 管制圖：模擬 25 批 × 50 顆晶粒，注入漂移情境，套用改善後 (After) 實測判定誤差。

真實缺陷 → 真實判定（依檢驗規範 v1 的嚴重度分級）→ 依一致性分析量到的
P(觀測判定 | 真實判定) 抽樣得到「系統觀測判定」→ p 管制圖（批不良率 = Fail 比例）。

Outputs
  results/control-chart.xlsx   資料、原生 Excel 折線圖、柏拉圖、參數
  results/p-chart.png, results/pareto.png
  public/data/lots.json        web app 讀取（schema 固定，不含管制界限）

Usage: python experiment/simulate_lots.py
"""
from __future__ import annotations

import json
import math
import sys
from pathlib import Path

import numpy as np

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
RES = HERE / "results"
CODES = ["CHP", "CRK", "SCR", "CON"]
LABELS = ["Pass", "Warning", "Fail"]

# ---------------------------------------------------------------- parameters (all simulated / assumed)
SEED = 20261005
N_LOTS, N_DIES = 25, 50
PHASE1_LOTS = 14          # lots 1..14 → baseline for control limits
DRIFT_START = 15          # 漂移情境 starts at lot 15
# per-die probability that a defect of this class is present, and its severity split
# (Pass-level = 門檻以下, Warning-level, Fail-level) per 檢驗規範 v1
BASE = {
    "CHP": {"p": 0.040, "sev": (0.20, 0.50, 0.30)},
    "CRK": {"p": 0.006, "sev": (0.00, 0.00, 1.00)},
    "SCR": {"p": 0.030, "sev": (0.20, 0.50, 0.30)},
    "CON": {"p": 0.050, "sev": (0.25, 0.50, 0.25)},
}
# 切割刀磨耗：CHP occurrence rises linearly each lot from lot 15, and chips get deeper
CHP_P_STEP = 0.035        # +3.5 percentage points of CHP occurrence per lot after drift start
CHP_FAIL_SHARE_END = 0.70 # Fail share of CHP rises linearly from 0.30 (lot 14) to this at lot 25
SCENARIO = "自第 15 批起切割刀磨耗，CHP 崩角比例逐批上升（模擬）"


def chp_params(lot):
    if lot < DRIFT_START:
        return BASE["CHP"]["p"], BASE["CHP"]["sev"]
    k = lot - (DRIFT_START - 1)
    p = BASE["CHP"]["p"] + CHP_P_STEP * k
    fail = 0.30 + (CHP_FAIL_SHARE_END - 0.30) * k / (N_LOTS - (DRIFT_START - 1))
    sub = 0.20
    return p, (sub, 1 - sub - fail, fail)


def load_error_model():
    """P(observed verdict | true verdict, class of the verdict-determining defect), measured in the
    After condition (analyze_msa.py → msa-metrics.json). Rows with no measured samples fall back to the
    pooled row for that true verdict. Unparsable AI output (Invalid) is routed to 人工複判 = Warning."""
    m = json.loads((RES / "msa-metrics.json").read_text(encoding="utf-8"))
    conf, by_class = m["afterConfusion"], m["afterConfusionByClass"]

    def vec(row):
        return np.array([row["Pass"], row["Warning"] + row.get("Invalid", 0), row["Fail"]], dtype=float)

    pooled = {t: vec(conf[t]) for t in LABELS}
    raw = {}
    for t in LABELS:
        for c in CODES + [None]:
            counts = vec(by_class[t][c]) if (t in by_class and c is not None) else pooled[t]
            if counts.sum() == 0:
                counts = pooled[t]
            raw[(t, c)] = counts / counts.sum() if counts.sum() else np.eye(3)[LABELS.index(t)]
    # hallucinated classes (AI false positives), used to label false-call dies in the Pareto
    fp = np.array([m["defectRecognition"][c]["fp"] for c in CODES], dtype=float) + 0.5
    rates = {"miss": m["metrics"]["after"]["missRate"] or 0.0, "falseCall": m["metrics"]["after"]["falseCallRate"] or 0.0}
    return raw, fp / fp.sum(), rates, conf


def apply_mode(raw, mode):
    """mode 'rejudge' (primary, proposed deployment): every AI-Warning die goes to 人工複判, assumed to
    reach the true verdict (true Pass → Pass, true Fail → Fail, true Warning stays Warning).
    mode 'raw' (sensitivity): the measured AI+rule verdict is final, no re-judgment."""
    if mode == "raw":
        return raw
    if mode == "perfect":  # reference only: judgement equals the true verdict
        return {k: np.eye(3)[LABELS.index(k[0])] for k in raw}
    out = {}
    for (t, c), p in raw.items():
        q = p.copy()
        w = q[1]
        q[1] = 0.0
        q[LABELS.index(t)] += w
        out[(t, c)] = q
    return out


def simulate(mode="rejudge"):
    raw, fp_mix, rates, conf = load_error_model()
    P = apply_mode(raw, mode)
    rng = np.random.default_rng(SEED)          # true defect process (identical across modes)
    rng_obs = np.random.default_rng(SEED + 1)  # judgement layer
    lots = []
    for lot in range(1, N_LOTS + 1):
        obs_counts = {"Pass": 0, "Warning": 0, "Fail": 0}
        true_counts = {"Pass": 0, "Warning": 0, "Fail": 0}
        defect_counts = {c: 0 for c in CODES}
        true_defect_counts = {c: 0 for c in CODES}
        for _ in range(N_DIES):
            sev_by_class = {}
            for c in CODES:
                p, sev = chp_params(lot) if c == "CHP" else (BASE[c]["p"], BASE[c]["sev"])
                if rng.random() < p:
                    sev_by_class[c] = int(rng.choice(3, p=sev))  # 0 Pass-level,1 Warning,2 Fail
            true_rank = max(sev_by_class.values(), default=0)
            true_v = LABELS[true_rank]
            worst = [c for c, s in sev_by_class.items() if s == true_rank and true_rank > 0]
            cls = worst[int(rng_obs.integers(len(worst)))] if worst else None
            obs_v = LABELS[int(rng_obs.choice(3, p=P[(true_v, cls)]))]
            true_counts[true_v] += 1
            obs_counts[obs_v] += 1
            relevant = [c for c, s in sev_by_class.items() if s >= 1]
            for c in relevant:
                true_defect_counts[c] += 1
            if obs_v != "Pass":
                if relevant:
                    for c in relevant:
                        defect_counts[c] += 1
                else:  # false call / false warning: label with the AI's typical hallucinated class
                    defect_counts[CODES[int(rng_obs.choice(4, p=fp_mix))]] += 1
        lots.append({"lotId": f"L{lot:02d}", "n": N_DIES, "pass": obs_counts["Pass"], "warning": obs_counts["Warning"],
                     "fail": obs_counts["Fail"], "defectCounts": defect_counts,
                     "_true": true_counts, "_trueDefects": true_defect_counts})
    return lots, rates, P, conf


def control_limits(lots):
    base = lots[:PHASE1_LOTS]
    pbar = sum(l["fail"] for l in base) / sum(l["n"] for l in base)
    sigma = math.sqrt(pbar * (1 - pbar) / N_DIES)
    return pbar, sigma


def signals(lots, pbar, sigma):
    ucl = pbar + 3 * sigma
    lcl = max(0.0, pbar - 3 * sigma)
    p = [l["fail"] / l["n"] for l in lots]
    out = []
    for i, x in enumerate(p):
        s = []
        if x > ucl or (lcl > 0 and x < lcl):
            s.append("異常：超出管制界限")
        if i >= 6 and all(p[j] > p[j - 1] for j in range(i - 5, i + 1)):
            s.append("預警：連續 7 點上升")
        if i >= 2:
            w = p[i - 2:i + 1]
            if sum(v > pbar + 2 * sigma for v in w) >= 2 and x > pbar + 2 * sigma:
                s.append("預警：3 點中 2 點超過 +2σ")
            if pbar - 2 * sigma > 0 and sum(v < pbar - 2 * sigma for v in w) >= 2 and x < pbar - 2 * sigma:
                s.append("預警：3 點中 2 點低於 −2σ")
        out.append(s)
    return p, ucl, lcl, out


def font_prop():
    from matplotlib import font_manager
    for f in (r"C:\Windows\Fonts\msjh.ttc", r"C:\Windows\Fonts\msyh.ttc"):
        if Path(f).exists():
            font_manager.fontManager.addfont(f)
            return font_manager.FontProperties(fname=f).get_name()
    return None


def plots(lots, p, pbar, sigma, ucl, sig, first_warn, first_ooc):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    name = font_prop()
    if name:
        plt.rcParams["font.family"] = name
    plt.rcParams["axes.unicode_minus"] = False
    x = np.arange(1, len(lots) + 1)
    fig, ax = plt.subplots(figsize=(11, 5.2), dpi=150)
    ax.axvspan(DRIFT_START - 0.5, len(lots) + 0.5, color="#f3e9dc", zorder=0, label="漂移情境（模擬）")
    ax.plot(x, p, "-o", color="#2f4f6f", lw=1.8, ms=5, label="批不良率 p（判定 Fail 比例）", zorder=3)
    ax.axhline(pbar, color="#555", lw=1.2, label=f"CL = {pbar:.3f}")
    ax.axhline(ucl, color="#c0392b", lw=1.2, ls="--", label=f"UCL (+3σ) = {ucl:.3f}")
    ax.axhline(pbar + 2 * sigma, color="#e08e45", lw=1, ls=":", label=f"+2σ = {pbar + 2 * sigma:.3f}")
    for i, s in enumerate(sig):
        if any(t.startswith("異常") for t in s):
            ax.plot(x[i], p[i], "o", ms=10, mfc="none", mec="#c0392b", mew=2, zorder=4)
        elif s:
            ax.plot(x[i], p[i], "s", ms=10, mfc="none", mec="#e08e45", mew=2, zorder=4)
    for i, s_ in enumerate(sig[:PHASE1_LOTS]):
        if any(t.startswith("異常") for t in s_):
            ax.annotate("Phase I 誤警報（基準期隨機波動）", (x[i], p[i]), xytext=(12, 8), textcoords="offset points",
                        fontsize=8, color="#c0392b")
    ax.text(PHASE1_LOTS / 2 + 0.5, 0.97, "Phase I：建立管制界限", transform=ax.get_xaxis_transform(), ha="center", fontsize=8, color="#555")
    ax.text((PHASE1_LOTS + 1 + len(lots)) / 2, 0.97, "Phase II：監控", transform=ax.get_xaxis_transform(), ha="center", fontsize=8, color="#555")
    if first_warn and first_warn == first_ooc:
        ax.annotate(f"首次預警＝首次異常 L{first_ooc:02d}（提早 0 批）", (first_ooc, p[first_ooc - 1]), xytext=(-190, 10),
                    textcoords="offset points", arrowprops=dict(arrowstyle="->", color="#c0392b"), color="#c0392b")
        first_warn = first_ooc = None
    if first_warn:
        ax.annotate(f"首次預警 L{first_warn:02d}", (first_warn, p[first_warn - 1]), xytext=(-70, 30),
                    textcoords="offset points", arrowprops=dict(arrowstyle="->", color="#e08e45"), color="#b9651f")
    if first_ooc:
        ax.annotate(f"首次異常 L{first_ooc:02d}", (first_ooc, p[first_ooc - 1]), xytext=(-80, 18),
                    textcoords="offset points", arrowprops=dict(arrowstyle="->", color="#c0392b"), color="#c0392b")
    ax.set_xticks(x)
    ax.set_xticklabels([l["lotId"] for l in lots], rotation=60, fontsize=8)
    ax.set_xlim(0.5, len(lots) + 0.5)
    ax.set_ylim(0, max(max(p), ucl) * 1.2)
    ax.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
    ax.set_ylabel("批不良率")
    ax.set_title(f"p 管制圖（每批 n={N_DIES}，管制界限以 L01–L{PHASE1_LOTS:02d} 計算；數據為模擬）")
    ax.grid(axis="y", alpha=0.3)
    ax.legend(loc="upper left", bbox_to_anchor=(0, 0.93), fontsize=8, ncol=2, frameon=False)
    fig.tight_layout()
    fig.savefig(RES / "p-chart.png")
    plt.close(fig)

    fig, axes = plt.subplots(1, 2, figsize=(11, 4.6), dpi=150, sharey=False)
    for ax, (title, sl) in zip(axes, [(f"L01–L{PHASE1_LOTS:02d}（漂移前）", lots[:PHASE1_LOTS]),
                                      (f"L{DRIFT_START:02d}–L{len(lots):02d}（漂移後）", lots[DRIFT_START - 1:])]):
        tot = {c: sum(l["defectCounts"][c] for l in sl) for c in CODES}
        order = sorted(CODES, key=lambda c: -tot[c])
        vals = [tot[c] for c in order]
        cum = np.cumsum(vals) / max(1, sum(vals))
        ax.bar(order, vals, color=["#c0392b" if c == "CHP" else "#7f8c9a" for c in order])
        for i, v in enumerate(vals):
            ax.text(i, v, str(v), ha="center", va="bottom", fontsize=9)
        ax2 = ax.twinx()
        ax2.plot(order, cum, "-o", color="#2f4f6f", ms=4)
        ax2.set_ylim(0, 1.05)
        ax2.yaxis.set_major_formatter(matplotlib.ticker.PercentFormatter(1.0))
        ax.set_title(title)
        ax.set_ylabel("缺陷次數（系統判定非 Pass 之晶粒）")
        ax.set_ylim(0, max(vals) * 1.25 if max(vals) else 1)
    fig.suptitle("缺陷分類柏拉圖：漂移前 vs 漂移後（數據為模擬）")
    fig.tight_layout()
    fig.savefig(RES / "pareto.png")
    plt.close(fig)


def excel(lots, p, pbar, sigma, ucl, lcl, sig, rates, raw, first_warn, first_ooc, sens):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "p管制圖"
    hdr = ["批", "n", "Pass", "Warning", "Fail", "批不良率 p", "CL", "UCL", "LCL", "+2σ", "判異訊號",
           "（真實）Fail", "（真實）Warning"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2F4F6F")
    for i, l in enumerate(lots):
        r = i + 2
        ws.append([l["lotId"], l["n"], l["pass"], l["warning"], l["fail"], f"=E{r}/B{r}", round(pbar, 5),
                   round(ucl, 5), round(lcl, 5), round(pbar + 2 * sigma, 5), "；".join(sig[i]),
                   l["_true"]["Fail"], l["_true"]["Warning"]])
        for col in "FGHIJ":
            ws[f"{col}{r}"].number_format = "0.0%"
        if sig[i]:
            ws[f"K{r}"].font = Font(color="C0392B" if any(s.startswith("異常") for s in sig[i]) else "B9651F", bold=True)
    for col, w in zip("ABCDEFGHIJKLM", (6, 5, 7, 9, 6, 12, 8, 8, 8, 8, 34, 12, 14)):
        ws.column_dimensions[col].width = w
    last = len(lots) + 1
    ch = LineChart()
    ch.title = "p 管制圖（模擬）"
    ch.y_axis.title = "批不良率"
    ch.x_axis.title = "批"
    ch.y_axis.number_format = "0%"
    ch.height, ch.width = 9, 22
    ch.add_data(Reference(ws, min_col=6, max_col=8, min_row=1, max_row=last), titles_from_data=True)
    ch.add_data(Reference(ws, min_col=10, max_col=10, min_row=1, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
    styles = [("2F4F6F", None), ("555555", None), ("C0392B", "dash"), ("E08E45", "sysDot")]
    for s, (color, dash) in zip(ch.series, styles):
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 20000
        if dash:
            s.graphicalProperties.line.dashStyle = dash
        s.smooth = False
    ch.series[0].marker.symbol = "circle"
    ws.add_chart(ch, "O2")
    r0 = last + 2
    ws.cell(row=r0, column=1, value="管制界限以 L01–L14（Phase I 基準期）計算；p̄ 與 σ = sqrt(p̄(1−p̄)/n)").font = Font(italic=True)
    ws.cell(row=r0 + 1, column=1, value=f"首次預警：{'L%02d' % first_warn if first_warn else '無'}；首次異常：{'L%02d' % first_ooc if first_ooc else '無'}")

    ps = wb.create_sheet("柏拉圖")
    ps.append(["缺陷分類", f"L01–L{PHASE1_LOTS:02d}", f"L{DRIFT_START:02d}–L{N_LOTS:02d}"])
    for c in CODES:
        ps.append([c, sum(l["defectCounts"][c] for l in lots[:PHASE1_LOTS]),
                   sum(l["defectCounts"][c] for l in lots[DRIFT_START - 1:])])
    bc = BarChart()
    bc.title = "缺陷分類次數：漂移前 vs 漂移後（模擬）"
    bc.add_data(Reference(ps, min_col=2, max_col=3, min_row=1, max_row=5), titles_from_data=True)
    bc.set_categories(Reference(ps, min_col=1, min_row=2, max_row=5))
    bc.height, bc.width = 8, 16
    ps.add_chart(bc, "E2")
    ps.append([])
    ps.append(["各批明細"])
    ps.append(["批"] + CODES)
    for l in lots:
        ps.append([l["lotId"]] + [l["defectCounts"][c] for c in CODES])

    qs = wb.create_sheet("模擬參數")
    lines = [
        "所有數據為模擬（非真實產線數據）。",
        f"亂數種子 {SEED}；{N_LOTS} 批 × 每批 {N_DIES} 顆。",
        "真實缺陷：每顆晶粒每類缺陷獨立以機率 p 出現，嚴重度依 (門檻以下, Warning 級, Fail 級) 比例抽樣；晶粒真實判定取最嚴重者。",
    ] + [f"  {c}: p = {BASE[c]['p']:.3f}，嚴重度比例 {BASE[c]['sev']}" for c in CODES] + [
        f"漂移情境：{SCENARIO}；自 L{DRIFT_START} 起 CHP 出現機率每批 +{CHP_P_STEP:.3f}，CHP 中 Fail 級比例由 0.30 線性升至 {CHP_FAIL_SHARE_END:.2f}（L{N_LOTS}）。",
        "判定層（主圖假設＝建議導入設計）：AI＋規則的判定依一致性分析「改善後」實測、按缺陷分類分開的混淆矩陣",
        "  P(AI 判定 | 真實判定, 決定判定的缺陷分類) 抽樣；AI 判 Warning 的晶粒一律送人工複判，假設人工複判得到真實判定。",
        "  AI 判 Pass（含漏判）直接放行、AI 判 Fail（含誤判）直接退件，不經人工 → 漏判率、誤判率維持改善後實測值。",
        "敏感度分析（另一工作表）：同一真實缺陷流程，但不做人工複判、AI＋規則判定即最終判定。",
    ] + [f"  真實 {t}／{c or '—'} → AI 判定 Pass/Warning/Fail = {', '.join(f'{v:.3f}' for v in raw[(t, c)])}"
         for t in LABELS for c in (CODES if t != "Pass" else [None])] + [
        f"漏判率 = {rates['miss']:.3f}；誤判率 = {rates['falseCall']:.3f}（改善後實測，lots.json judgementErrorRates）",
        "判異規則：點超出 UCL → 異常；連續 7 點上升、或 3 點中 2 點超過 +2σ（同側）→ 預警。",
        "柏拉圖：只計系統判定非 Pass 的晶粒上之 Warning／Fail 級缺陷；誤判晶粒依 AI 多報缺陷分類比例標記代碼。",
    ]
    for line in lines:
        qs.append([line])
    qs.column_dimensions["A"].width = 120

    ss = wb.create_sheet("敏感度分析")
    ss.append(["判定層：無人工複判（AI＋規則判定即最終判定，改善後實測誤差）；真實缺陷流程與主圖相同"])
    ss.append([f"CL = {sens['pbar']:.4f}；UCL = {sens['ucl']:.4f}；首次預警 = {sens['firstWarningLot']}；首次異常 = {sens['firstOutOfControlLot']}"])
    ss.append(["批", "n", "Pass", "Warning", "Fail", "批不良率 p", "判異訊號"])
    for l, x, sg in zip(sens["lots"], sens["p"], sens["signals"]):
        ss.append([l["lotId"], l["n"], l["pass"], l["warning"], l["fail"], x, "；".join(sg)])
        ss.cell(row=ss.max_row, column=6).number_format = "0.0%"
    ss.column_dimensions["G"].width = 34
    wb.save(RES / "control-chart.xlsx")


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    def run(mode):
        lots, rates, raw, conf = simulate(mode)
        pbar, sigma = control_limits(lots)
        p, ucl, lcl, sig = signals(lots, pbar, sigma)
        # 提早 N 批 is measured in Phase II (monitoring, lots after the Phase I baseline); Phase I signals are reported separately
        first_ooc = next((i + 1 for i, s in enumerate(sig) if i >= PHASE1_LOTS and any(t.startswith("異常") for t in s)), None)
        first_warn = next((i + 1 for i, s in enumerate(sig) if i >= PHASE1_LOTS and any(t.startswith("預警") for t in s)), None)
        return lots, rates, raw, pbar, sigma, p, ucl, lcl, sig, first_warn, first_ooc

    (sl, _, _, s_pbar, _, s_p, s_ucl, _, s_sig, s_fw, s_fo) = run("raw")
    sens = {"lots": sl, "pbar": s_pbar, "ucl": s_ucl, "p": s_p, "signals": s_sig,
            "firstWarningLot": s_fw, "firstOutOfControlLot": s_fo,
            "leadLots": (s_fo - s_fw) if (s_fw and s_fo) else None}
    print(f"[sensitivity: no 人工複判] p̄={s_pbar:.4f} UCL={s_ucl:.4f} first 預警={s_fw} first 異常={s_fo}")
    print("  fails per lot:", [l["fail"] for l in sl])
    (_, _, _, pf_pbar, _, pf_p, pf_ucl, _, pf_sig, pf_fw, pf_fo) = run("perfect")
    perfect = {"pbar": pf_pbar, "ucl": pf_ucl, "firstWarningLot": pf_fw, "firstOutOfControlLot": pf_fo,
               "leadLots": (pf_fo - pf_fw) if (pf_fw and pf_fo) else None}
    print(f"[reference: perfect judgement] p̄={pf_pbar:.4f} UCL={pf_ucl:.4f} first 預警={pf_fw} first 異常={pf_fo}")
    lots, rates, raw, pbar, sigma, p, ucl, lcl, sig, first_warn, first_ooc = run("rejudge")
    plots(lots, p, pbar, sigma, ucl, sig, first_warn, first_ooc)
    excel(lots, p, pbar, sigma, ucl, lcl, sig, rates, raw, first_warn, first_ooc, sens)
    out = {
        "specVersion": "v1",
        "scenario": SCENARIO,
        "simulated": True,
        "judgementErrorRates": {"miss": round(rates["miss"], 4), "falseCall": round(rates["falseCall"], 4)},
        "lots": [{k: l[k] for k in ("lotId", "n", "pass", "warning", "fail", "defectCounts")} for l in lots],
    }
    dst = REPO / "public" / "data" / "lots.json"
    dst.parent.mkdir(parents=True, exist_ok=True)
    dst.write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"p̄={pbar:.4f} σ={sigma:.4f} UCL={ucl:.4f} +2σ={pbar + 2 * sigma:.4f}")
    for l, x, s in zip(lots, p, sig):
        print(l["lotId"], f"obs P/W/F={l['pass']}/{l['warning']}/{l['fail']}", f"p={x:.2f}",
              f"true W/F={l['_true']['Warning']}/{l['_true']['Fail']}", l["defectCounts"], "；".join(s))
    lead = (first_ooc - first_warn) if (first_warn and first_ooc) else None
    print(f"first 預警 = {first_warn}, first 異常 = {first_ooc}, lead = {lead}")
    summary = {"pbar": pbar, "sigma": sigma, "ucl": ucl, "firstWarningLot": first_warn,
               "firstOutOfControlLot": first_ooc, "leadLots": lead, "signals": sig, "seed": SEED,
               "judgementLayer": "measured After per-class errors + 人工複判 of AI-Warning dies (assumed correct)",
               "sensitivityNoRejudge": {k: v for k, v in sens.items() if k != "lots"},
               "referencePerfectJudgement": perfect,
               "phaseISignals": {f"L{i + 1:02d}": s_ for i, s_ in enumerate(sig[:PHASE1_LOTS]) if s_},
               "trueFailPerLot": [l["_true"]["Fail"] for l in lots]}
    (RES / "control-chart-summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
