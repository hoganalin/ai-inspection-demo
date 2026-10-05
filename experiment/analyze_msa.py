"""Step 2b — 一致性分析計算與輸出。

Reads experiment/raw/{before,after}.jsonl (written by run_msa.py), derives verdicts
(After: spec_v1.derive_verdict on the AI defect list), computes attribute-agreement
metrics and writes:
  results/msa-results.xlsx   (摘要, 明細, 各評估者, 缺陷辨識, 混淆矩陣, 實驗設計)
  results/msa-metrics.json   (machine-readable; used by msa-summary.md and simulate_lots.py)

Usage: python experiment/analyze_msa.py [--trials 2]
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import spec_v1 as spec
from run_msa import APPRAISERS, CONDITIONS, MODEL, PERSONA, VERBAL, load_cache, parse_after, parse_before, spent

HERE = Path(__file__).resolve().parent
RES = HERE / "results"
LABELS = ["Pass", "Warning", "Fail"]
INVALID = "Invalid"
COND_ZH = {"before": "改善前 (口述標準)", "after": "改善後 (規範 v1)"}
MEAS_KEY = {"CHP": "depthUm", "SCR": "lengthUm", "CRK": "lengthUm", "CON": "diameterUm"}


def fleiss_kappa(table):
    """table: list of per-item category counts (each row sums to n raters)."""
    N = len(table)
    if N == 0:
        return None
    n = sum(table[0])
    k = len(table[0])
    p_j = [sum(r[j] for r in table) / (N * n) for j in range(k)]
    P_i = [(sum(c * c for c in r) - n) / (n * (n - 1)) for r in table]
    P_bar = sum(P_i) / N
    P_e = sum(p * p for p in p_j)
    return None if P_e == 1 else (P_bar - P_e) / (1 - P_e)


def is_relevant(d):
    """Defect that can change the verdict by itself (not sub-threshold)."""
    if d["code"] == "CON" and d["zone"] == "core":
        return d.get("diameterUm", 0) >= 20
    return spec.judge_defect(d) != spec.PASS


def match_defects(ref, pred):
    """Per-class greedy matching by measurement rank. Returns (pairs, fn_list, fp_list)."""
    pairs, fns, fps = [], [], []
    for code in spec.CODES:
        r = sorted([d for d in ref if d["code"] == code], key=lambda d: -d.get(MEAS_KEY[code], 0))
        p = sorted([d for d in pred if d.get("code") == code],
                   key=lambda d: -float(d.get(MEAS_KEY[code]) or 0))
        m = min(len(r), len(p))
        pairs += list(zip(r[:m], p[:m]))
        fns += r[m:]
        fps += p[m:]
    fps += [d for d in pred if d.get("code") not in spec.CODES]
    return pairs, fns, fps


def collect(trials):
    ref = json.loads((HERE / "samples" / "reference.json").read_text(encoding="utf-8"))
    samples = {s["id"]: s for s in ref["samples"]}
    cache = load_cache()
    rows = []
    for rec in cache.values():
        if rec["trial"] > trials or rec["sample"] not in samples:
            continue
        s = samples[rec["sample"]]
        if rec["condition"] == "before":
            v = parse_before(rec["text"]) or INVALID
            defects = None
        else:
            defects = parse_after(rec["text"])
            v = spec.derive_verdict(defects)["verdict"] if defects is not None else INVALID
        rows.append({**{k: rec[k] for k in ("condition", "appraiser", "sample", "trial", "costUsd")},
                     "reference": s["referenceVerdict"], "verdict": v, "aiDefects": defects,
                     "refDefects": s["defects"], "text": rec["text"]})
    return samples, rows


def complete_ids(rows, cond, trials):
    have = defaultdict(set)
    for r in rows:
        if r["condition"] == cond:
            have[r["sample"]].add((r["appraiser"], r["trial"]))
    need = {(a, t) for a in APPRAISERS for t in range(1, trials + 1)}
    return sorted(i for i, s in have.items() if need <= s)


def metrics(rows, cond, ids, samples, trials, all_rows):
    V = {(r["appraiser"], r["sample"], r["trial"]): r["verdict"] for r in rows if r["condition"] == cond}
    VA = {(r["appraiser"], r["sample"], r["trial"]): r["verdict"] for r in all_rows if r["condition"] == cond}
    out = {"n_images": len(ids), "trials": trials, "per_appraiser": {}}
    T = range(1, trials + 1)
    js_all = []
    rep_agree = rep_n = 0
    for a in APPRAISERS:
        # 重複性: images where this appraiser has >= 2 trials (may be fewer than ids if a run was cut short)
        pairs = [i for i in ids if (a, i, 1) in VA and (a, i, 2) in VA]
        agree = sum(len({VA[(a, i, t)] for t in (1, 2, 3) if (a, i, t) in VA}) == 1 for i in pairs)
        rep = agree / len(pairs) if pairs else None
        rep_agree += agree
        rep_n += len(pairs)
        vs_std = sum(all(V[(a, i, t)] == samples[i]["referenceVerdict"] for t in T) for i in ids) / len(ids)
        js = [(samples[i]["referenceVerdict"], V[(a, i, t)]) for i in ids for t in T]
        js_all += js
        out["per_appraiser"][a] = {"repeatability": rep, "repeatabilityN": len(pairs), "vsStandardAllTrials": vs_std,
                                   **judgment_rates(js)}
    out["repeatability"] = rep_agree / rep_n if rep_n else None
    out["repeatabilityN"] = rep_n
    out["reproducibilityAll"] = sum(len({V[(a, i, t)] for a in APPRAISERS for t in T}) == 1 for i in ids) / len(ids)
    out["betweenAppraisersPerTrial"] = sum(
        len({V[(a, i, t)] for a in APPRAISERS}) == 1 for i in ids for t in T) / (len(ids) * trials)
    out["allVsStandard"] = sum(all(V[(a, i, t)] == samples[i]["referenceVerdict"] for a in APPRAISERS for t in T)
                               for i in ids) / len(ids)
    out.update(judgment_rates(js_all))
    cats = LABELS + [INVALID]
    table = [[sum(V[(a, i, t)] == c for a in APPRAISERS for t in T) for c in cats] for i in ids]
    out["fleissKappa"] = fleiss_kappa(table)
    conf = {r_: {c: 0 for c in cats} for r_ in LABELS}
    for r_, v in js_all:
        conf[r_][v] += 1
    out["confusion"] = conf
    out["invalidCount"] = sum(v == INVALID for _, v in js_all)
    return out


def judgment_rates(js):
    n = len(js)
    fails = [v for r, v in js if r == "Fail"]
    passes = [v for r, v in js if r == "Pass"]
    return {
        "judgments": n,
        "accuracy": sum(r == v for r, v in js) / n,
        "missRate": (sum(v == "Pass" for v in fails) / len(fails)) if fails else None,
        "falseCallRate": (sum(v == "Fail" for v in passes) / len(passes)) if passes else None,
        "warningRate": sum(v == "Warning" for _, v in js) / n,
        "referenceWarningRate": sum(r == "Warning" for r, _ in js) / n,
    }


def defect_recognition(rows, ids):
    per = {c: Counter() for c in spec.CODES}
    errs = defaultdict(list)
    halluc_clean = Counter()
    clean_judgments = 0
    zone_ok = Counter()
    for r in rows:
        if r["condition"] != "after" or r["sample"] not in ids or r["aiDefects"] is None:
            continue
        pairs, fns, fps = match_defects(r["refDefects"], r["aiDefects"])
        for rd, pd in pairs:
            c = rd["code"]
            per[c]["tp"] += 1
            per[c]["tp_rel" if is_relevant(rd) else "tp_sub"] += 1
            try:
                pv = float(pd.get(MEAS_KEY[c]) or 0)
            except (TypeError, ValueError):
                pv = 0.0
            errs[c].append((rd.get(MEAS_KEY[c], 0), pv))
            zone_ok[c, pd.get("zone") == rd["zone"]] += 1
        for rd in fns:
            per[rd["code"]]["fn"] += 1
            per[rd["code"]]["fn_rel" if is_relevant(rd) else "fn_sub"] += 1
        for pd in fps:
            c = pd.get("code") if pd.get("code") in spec.CODES else "CON"
            per[c]["fp"] += 1
        if not r["refDefects"]:
            clean_judgments += 1
            for pd in r["aiDefects"]:
                halluc_clean[pd.get("code")] += 1
    out = {}
    for c in spec.CODES:
        p = per[c]
        e = errs[c]
        mae = sum(abs(pv - rv) for rv, pv in e) / len(e) if e else None
        bias = sum(pv - rv for rv, pv in e) / len(e) if e else None
        rel = sum((pv - rv) / rv for rv, pv in e if rv) / len(e) if e else None
        out[c] = {
            "refInstances": p["tp"] + p["fn"], "tp": p["tp"], "fn": p["fn"], "fp": p["fp"],
            "recall": p["tp"] / (p["tp"] + p["fn"]) if p["tp"] + p["fn"] else None,
            "recallRelevant": p["tp_rel"] / (p["tp_rel"] + p["fn_rel"]) if p["tp_rel"] + p["fn_rel"] else None,
            "recallSubThreshold": p["tp_sub"] / (p["tp_sub"] + p["fn_sub"]) if p["tp_sub"] + p["fn_sub"] else None,
            "precision": p["tp"] / (p["tp"] + p["fp"]) if p["tp"] + p["fp"] else None,
            "measMaeUm": mae, "measBiasUm": bias, "measRelBias": rel,
            "zoneAccuracy": zone_ok[c, True] / (zone_ok[c, True] + zone_ok[c, False]) if e else None,
            "pairs": e,
        }
    out["_clean"] = {"judgments": clean_judgments, "hallucinatedByCode": dict(halluc_clean)}
    return out


# ---------------------------------------------------------------- Excel
HDR = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="2F4F6F")


def header(ws, row, values):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def autosize(ws, widths=None):
    for j, col in enumerate(ws.columns, 1):
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(j)].width = (widths or {}).get(j, min(60, max(10, w * 1.6)))


def pct(cell):
    cell.number_format = "0.0%"


SUMMARY_ROWS = [
    ("重複性（同一評估者兩次判定一致的比例；只計有兩次判定的樣本）", "repeatability", True),
    ("重複性樣本數（評估者×樣本配對數）", "repeatabilityN", False),
    ("再現性（全部評估者 × 全部試次判定皆一致的樣本比例）", "reproducibilityAll", True),
    ("評估者間一致率（同一試次三位評估者一致，平均）", "betweenAppraisersPerTrial", True),
    ("對標準答案一致率（逐筆判定）", "accuracy", True),
    ("對標準答案一致率（全部評估者全部試次皆對的樣本比例）", "allVsStandard", True),
    ("漏判率（標準答案 Fail 被判 Pass）", "missRate", True),
    ("誤判率（標準答案 Pass 被判 Fail）", "falseCallRate", True),
    ("人工複判率（判定 Warning 的比例）", "warningRate", True),
    ("人工複判率參考值（標準答案中 Warning 比例）", "referenceWarningRate", True),
    ("Fleiss' kappa（全部評估者×試次，3 類＋無效）", "fleissKappa", False),
    ("無效回覆數（無法解析）", "invalidCount", False),
    ("判定筆數", "judgments", False),
]


def write_excel(M, rows, ids_by_cond, samples, defrec, design):
    wb = Workbook()
    ws = wb.active
    ws.title = "摘要"
    ws["A1"] = "一致性分析摘要：改善前 (口述標準) vs 改善後 (檢驗規範 v1)"
    ws["A1"].font = Font(bold=True, size=13)
    ws["A2"] = f"模型 {MODEL}；評估者 3 位（模擬 AI 評估者）；樣本 {design['nImages']} 張；每組 {design['trials']} 次；樣本與評估者皆為模擬"
    header(ws, 4, ["指標", COND_ZH["before"], COND_ZH["after"], "差異 (後−前)"])
    for k, (label, key, is_pct) in enumerate(SUMMARY_ROWS, 5):
        ws.cell(row=k, column=1, value=label)
        b, a = M["before"].get(key), M["after"].get(key)
        ws.cell(row=k, column=2, value=b)
        ws.cell(row=k, column=3, value=a)
        if isinstance(a, (int, float)) and isinstance(b, (int, float)):
            ws.cell(row=k, column=4, value=a - b)
        if is_pct:
            for j in (2, 3, 4):
                pct(ws.cell(row=k, column=j))
        else:
            for j in (2, 3, 4):
                ws.cell(row=k, column=j).number_format = "0.000"
    autosize(ws, {1: 58, 2: 20, 3: 20, 4: 16})

    ws = wb.create_sheet("明細")
    header(ws, 1, ["條件", "評估者", "樣本", "試次", "標準答案", "判定", "與標準答案一致", "AI 缺陷清單 (After)", "標準答案缺陷", "原始回覆"])
    for r in sorted(rows, key=lambda r: (r["condition"] != "before", r["appraiser"], r["sample"], r["trial"])):
        ws.append([COND_ZH[r["condition"]], r["appraiser"], r["sample"], r["trial"], r["reference"], r["verdict"],
                   "是" if r["verdict"] == r["reference"] else "否",
                   json.dumps(r["aiDefects"], ensure_ascii=False) if r["aiDefects"] is not None else "",
                   json.dumps([{k: v for k, v in d.items() if k != "location"} for d in r["refDefects"]], ensure_ascii=False),
                   r["text"]])
    autosize(ws, {8: 50, 9: 50, 10: 80})
    ws.freeze_panes = "A2"

    ws = wb.create_sheet("各評估者")
    header(ws, 1, ["條件", "評估者", "標準 / 角色設定", "重複性", "重複性配對數", "對標準答案一致率（逐筆）", "全部試次皆對的樣本比例",
                   "漏判率", "誤判率", "人工複判率", "判定筆數"])
    for cond in CONDITIONS:
        for a in APPRAISERS:
            p = M[cond]["per_appraiser"][a]
            ws.append([COND_ZH[cond], a, VERBAL[a] if cond == "before" else PERSONA[a] + "（規範 v1）",
                       p["repeatability"], p["repeatabilityN"], p["accuracy"], p["vsStandardAllTrials"], p["missRate"],
                       p["falseCallRate"], p["warningRate"], p["judgments"]])
            for j in (4, 6, 7, 8, 9, 10):
                pct(ws.cell(row=ws.max_row, column=j))
    autosize(ws, {3: 44})

    ws = wb.create_sheet("缺陷辨識")
    ws["A1"] = "改善後 (After)：AI 回報缺陷 vs 標準答案缺陷（依缺陷分類、以量測值排序配對）"
    ws["A1"].font = Font(bold=True)
    header(ws, 3, ["缺陷分類", "標準答案缺陷次數", "偵出 (TP)", "漏報 (FN)", "多報 (FP)", "偵出率",
                   "偵出率（影響判定者）", "偵出率（門檻以下）", "精確率", "量測平均絕對誤差 (µm)",
                   "量測平均偏差 (µm)", "量測相對偏差", "區域判讀正確率"])
    for c in spec.CODES:
        d = defrec[c]
        ws.append([c, d["refInstances"], d["tp"], d["fn"], d["fp"], d["recall"], d["recallRelevant"],
                   d["recallSubThreshold"], d["precision"], d["measMaeUm"], d["measBiasUm"], d["measRelBias"],
                   d["zoneAccuracy"]])
        for j in (6, 7, 8, 9, 12, 13):
            pct(ws.cell(row=ws.max_row, column=j))
        for j in (10, 11):
            ws.cell(row=ws.max_row, column=j).number_format = "0.0"
    cl = defrec["_clean"]
    ws.append([])
    ws.append(["無缺陷樣本上的多報（幻覺）", f"{cl['judgments']} 筆判定", json.dumps(cl["hallucinatedByCode"], ensure_ascii=False)])
    ws.append([])
    header(ws, ws.max_row + 1, ["缺陷分類", "標準答案量測 (µm)", "AI 量測 (µm)"])
    for c in spec.CODES:
        for rv, pv in defrec[c]["pairs"]:
            ws.append([c, rv, pv])
    autosize(ws)

    ws = wb.create_sheet("混淆矩陣")
    r0 = 1
    for cond in CONDITIONS:
        ws.cell(row=r0, column=1, value=f"{COND_ZH[cond]}：列＝標準答案，欄＝評估者判定（筆數）").font = Font(bold=True)
        header(ws, r0 + 1, ["標準答案 \\ 判定"] + LABELS + [INVALID, "合計"])
        for k, rl in enumerate(LABELS):
            row = M[cond]["confusion"][rl]
            vals = [row[c] for c in LABELS + [INVALID]]
            for j, v in enumerate([rl] + vals + [sum(vals)], 1):
                ws.cell(row=r0 + 2 + k, column=j, value=v)
        r0 += 7
    autosize(ws)

    ws = wb.create_sheet("實驗設計")
    for line in design["lines"]:
        ws.append([line])
    ws.column_dimensions["A"].width = 120
    RES.mkdir(exist_ok=True)
    wb.save(RES / "msa-results.xlsx")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--trials", type=int, default=2)
    args = ap.parse_args()
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    samples, all_rows = collect(99)
    rows = [r for r in all_rows if r["trial"] <= args.trials]
    ids_by = {c: complete_ids(rows, c, args.trials) for c in CONDITIONS}
    ids = sorted(set(ids_by["before"]) & set(ids_by["after"]))
    if not ids:
        raise SystemExit("no complete samples yet")
    rows = [r for r in rows if r["sample"] in ids]
    M = {c: metrics(rows, c, ids, samples, args.trials, all_rows) for c in CONDITIONS}
    defrec = defect_recognition(rows, set(ids))
    total, ncalls = spent(load_cache())
    ref_mix = Counter(samples[i]["referenceVerdict"] for i in ids)
    design = {
        "nImages": len(ids), "trials": args.trials,
        "lines": [
            "實驗設計（皆為模擬）",
            f"模型：{MODEL}，temperature 1.0；每次呼叫獨立（無對話記憶）。",
            f"樣本：{len(ids)} 張半合成晶粒影像（程序化底圖＋注入已知尺寸缺陷），標準答案判定分布 {dict(ref_mix)}。",
            "評估者：甲／乙／丙 三位模擬 AI 評估者。",
            "改善前：三人 system prompt 完全相同，僅口述標準不同（甲 嚴：「" + VERBAL["甲"] + "」；乙 模糊：「" + VERBAL["乙"] + "」；丙 寬鬆：「" + VERBAL["丙"] + "」），AI 直接輸出 pass/warning/fail。",
            "改善後：三人共用同一份檢驗規範 v1 全文＋幾何與比例尺，只差一行中性角色（日班／夜班／假日班）；AI 只回報缺陷清單，判定由 spec_v1.py 依允收標準推導。",
            f"一致率、漏判率、誤判率、人工複判率、kappa 以第 1～{args.trials} 次判定計算（每位評估者 × 每條件 × 每張樣本）。",
            "原訂每組 2 次（預算 US$5 下由 3 次減為 2 次）；執行中 API 帳戶額度耗盡（HTTP 400 credit balance too low），"
            "第 2 次只完成改善前的部分判定，改善後 0 筆 → 重複性只能就有兩次判定的配對計算（見「重複性樣本數」），改善後重複性未量測。",
            "補齊方式：帳戶加值後執行 `run_msa.py run --trials 2 --budget 4.5`（自動從快取續跑，預估再花約 US$1.2），再以 --trials 2 重跑本分析。",
            f"API 總呼叫數（含前導測試）：{ncalls}；實際花費（依 usage × 牌價）US${total:.2f}。",
        ],
    }
    write_excel(M, [r for r in all_rows if r["sample"] in ids], ids_by, samples, defrec, design)
    after_conf = M["after"]["confusion"]
    # per-class confusion: key = class of the defect that determines the sample's reference verdict
    by_class = {t: {c: {k: 0 for k in LABELS + [INVALID]} for c in spec.CODES} for t in ("Warning", "Fail")}
    for r in rows:
        if r["condition"] != "after" or r["reference"] == "Pass":
            continue
        s = samples[r["sample"]]
        det = [d["code"] for d in s["defects"] if spec.judge_defect(d) == s["referenceVerdict"]]
        cls = det[0] if det else "CON"  # CON core point-count rule
        by_class[r["reference"]][cls][r["verdict"]] += 1
    out = {"design": design, "metrics": M, "defectRecognition": {k: {kk: vv for kk, vv in v.items() if kk != "pairs"}
                                                                   for k, v in defrec.items()},
           "afterConfusion": after_conf, "afterConfusionByClass": by_class, "spendUsd": total, "apiCalls": ncalls}
    (RES / "msa-metrics.json").write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8")
    # console summary
    for label, key, is_pct in SUMMARY_ROWS:
        b, a = M["before"].get(key), M["after"].get(key)
        fmt = (lambda x: "—" if x is None else f"{x:.1%}") if is_pct else (lambda x: "—" if x is None else f"{x:.3f}" if isinstance(x, float) else str(x))
        print(f"{label}: {fmt(b)} → {fmt(a)}")
    for cond in CONDITIONS:
        for a in APPRAISERS:
            p = M[cond]["per_appraiser"][a]
            print(cond, a, {k: round(v, 3) if isinstance(v, float) else v for k, v in p.items()})
        print(cond, "confusion", M[cond]["confusion"])
    for c in spec.CODES:
        print(c, {k: (round(v, 3) if isinstance(v, float) else v) for k, v in defrec[c].items() if k != "pairs"})
    print("clean", defrec["_clean"])
    print(f"images={len(ids)} spend=${total:.4f} calls={ncalls}")


if __name__ == "__main__":
    main()
