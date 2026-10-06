"""Export the reference sample set + recorded MSA replies for the web app.

Writes src/features/samples/referenceData.ts (generated, do not edit by hand):
  - 40 reference samples (標準答案 defects with locations, reference verdict)
  - the after-condition (檢驗規範 v1) AI replies of trial 1, per appraiser,
    parsed with the same rule as analyze_msa.py (last JSON object carrying a defects list)
  - the headline MSA metrics from results/msa-metrics.json

Usage:  python experiment/export_app_data.py
The web app derives verdicts from these defect lists with its own rule engine
(judgeDie); `npm run check:spec` verifies that it agrees with spec_v1.py.
"""
from __future__ import annotations

import json
from pathlib import Path

import spec_v1 as spec
from run_msa import load_cache, parse_after

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
OUT = REPO / "src" / "features" / "samples" / "referenceData.ts"

SHIFT = {"甲": "日班", "乙": "夜班", "丙": "假日班"}
DEFECT_KEYS = ("code", "zone", "depthUm", "lengthUm", "diameterUm", "crossesPad", "touchesSealRing")


def clean_ref(d: dict) -> dict:
    out = {k: d[k] for k in DEFECT_KEYS if k in d and d[k] is not None}
    if "location" in d:
        out["location"] = {"x": d["location"]["x"], "y": d["location"]["y"]}
    return out


HUMAN_FORM = HERE / "results" / "self-judgment-filled.xlsx"
HUMAN_META = HERE / "results" / "self-judgment-meta.json"
MEAS_KEY = {"CHP": "depthUm", "SCR": "lengthUm", "CON": "diameterUm"}


def _split(v) -> list[str]:
    return [x.strip() for x in str(v).replace("；", ";").split(";")] if v not in (None, "") else []


def human_defects(row) -> list[dict]:
    """表單一列（缺陷代碼、區域、量測）→ 缺陷清單；多個缺陷以 ; 對應。"""
    codes = [c for c in _split(row[2]) if c and c != "無"]
    zones, meas = _split(row[3]), _split(row[4])
    out = []
    for i, code in enumerate(codes):
        d = {"code": code, "zone": zones[i] if i < len(zones) else (zones[0] if zones else "core")}
        m = meas[i] if i < len(meas) else (meas[0] if len(meas) == 1 else None)
        if code in MEAS_KEY and m not in (None, ""):
            d[MEAS_KEY[code]] = float(m)
        out.append(d)
    return out


def rates(pairs: list[tuple[str, str]]) -> dict:
    """pairs = [(標準答案, 判定)]，判定用小寫。"""
    n = len(pairs)
    fails = [v for r, v in pairs if r == "fail"]
    passes = [v for r, v in pairs if r == "pass"]
    return {
        "n": n,
        "accuracy": sum(r == v for r, v in pairs) / n,
        "missRate": sum(v == "pass" for v in fails) / len(fails),
        "falseCallRate": sum(v == "fail" for v in passes) / len(passes),
        "warningRate": sum(v == "warning" for _, v in pairs) / n,
    }


def human_summary(ref_by_id: dict) -> tuple[dict | None, list]:
    """作者本人兩輪判定：自己判 vs 人眼回報＋規則推導。表單不存在時回傳 (None, [])。"""
    if not HUMAN_FORM.exists():
        return None, []
    from openpyxl import load_workbook  # 只有匯出人工結果時需要
    wb = load_workbook(HUMAN_FORM, data_only=True)
    meta = json.loads(HUMAN_META.read_text(encoding="utf-8"))
    replies, per_round = [], []
    for k, sheet in enumerate(("第1輪", "第2輪"), 1):
        own, rule = {}, {}
        for row in wb[sheet].iter_rows(min_row=2, max_row=41, values_only=True):
            verdict = str(row[5] or "").strip()
            if verdict not in ("Pass", "Warning", "Fail"):
                continue
            defects = human_defects(row)
            own[row[1]] = verdict.lower()
            rule[row[1]] = spec.derive_verdict(defects)["verdict"].lower()
            replies.append({"sample": row[1], "round": k, "defects": defects,
                            "ownVerdict": own[row[1]], "pyRuleVerdict": rule[row[1]]})
        refv = {s: ref_by_id[s]["referenceVerdict"].lower() for s in own}
        per_round.append({
            "round": k,
            "own": rates([(refv[s], own[s]) for s in own]),
            "rules": rates([(refv[s], rule[s]) for s in own]),
            "fixedByRules": sorted(s for s in own if own[s] != refv[s] and rule[s] == refv[s]),
            "brokenByRules": sorted(s for s in own if own[s] == refv[s] and rule[s] != refv[s]),
            "_own": own, "_rule": rule,
        })
    o1, o2 = per_round[0].pop("_own"), per_round[1].pop("_own")
    r1, r2 = per_round[0].pop("_rule"), per_round[1].pop("_rule")
    both = [s for s in o1 if s in o2]
    allp = [(ref_by_id[x["sample"]]["referenceVerdict"].lower(), x) for x in replies]
    minutes = meta["roundMinutes"]
    summary = {
        "appraiser": meta["appraiser"], "date": meta["date"], "note": meta["note"],
        "roundMinutes": minutes,
        "secondsPerDie": round(sum(minutes) * 60 / len(replies), 1),
        "rounds": per_round,
        "combined": {
            "own": rates([(r, x["ownVerdict"]) for r, x in allp]),
            "rules": rates([(r, x["pyRuleVerdict"]) for r, x in allp]),
            "fixedByRules": sum(len(p["fixedByRules"]) for p in per_round),
            "brokenByRules": sum(len(p["brokenByRules"]) for p in per_round),
        },
        "repeatability": {
            "n": len(both),
            "own": sum(o1[s] == o2[s] for s in both) / len(both),
            "rules": sum(r1[s] == r2[s] for s in both) / len(both),
        },
    }
    return summary, replies


def main() -> None:
    ref = json.loads((HERE / "samples" / "reference.json").read_text(encoding="utf-8"))
    metrics = json.loads((HERE / "results" / "msa-metrics.json").read_text(encoding="utf-8"))
    cache = load_cache()

    replies: dict[str, list] = {}
    for rec in cache.values():
        if rec["condition"] != "after" or rec["trial"] != 1:
            continue
        defects = parse_after(rec["text"])
        py_verdict = spec.derive_verdict(defects)["verdict"] if defects is not None else None
        replies.setdefault(rec["sample"], []).append({
            "appraiser": rec["appraiser"],
            "shift": SHIFT[rec["appraiser"]],
            "model": rec["model"],
            "defects": defects,
            "pyVerdict": py_verdict.lower() if py_verdict else None,
        })

    samples = []
    for s in ref["samples"]:
        rs = sorted(replies.get(s["id"], []), key=lambda r: "甲乙丙".index(r["appraiser"]))
        samples.append({
            "id": s["id"],
            "group": s["group"],
            "note": s["note"],
            "referenceVerdict": s["referenceVerdict"].lower(),
            "defects": [clean_ref(d) for d in s["defects"]],
            "recorded": rs,
        })

    keep = ("repeatability", "repeatabilityN", "betweenAppraisersPerTrial", "allVsStandard", "accuracy",
            "missRate", "falseCallRate", "warningRate", "referenceWarningRate", "fleissKappa", "judgments")
    msa = {
        "before": {k: metrics["metrics"]["before"].get(k) for k in keep},
        "after": {k: metrics["metrics"]["after"].get(k) for k in keep},
        "recognition": {
            code: {k: v.get(k) for k in ("refInstances", "tp", "fn", "recall", "measBiasUm")}
            for code, v in metrics["defectRecognition"].items() if code in ("CHP", "CRK", "SCR", "CON")
        },
    }

    human, human_replies = human_summary({s["id"]: s for s in ref["samples"]})

    body = (
        "// GENERATED by experiment/export_app_data.py — do not edit by hand.\n"
        "// Sources: experiment/samples/reference.json, experiment/raw/after.jsonl (trial 1),\n"
        "//          experiment/results/msa-metrics.json, experiment/results/self-judgment-filled.xlsx (+ meta).\n"
        "// Samples are semi-synthetic and AI appraisers are simulated; the human rounds are the author's own judgments.\n"
        "import type { RawReferenceSample, RawMsaSummary, RawHumanSummary, RawHumanReply } from './types';\n\n"
        f"export const GEOMETRY = {json.dumps(ref['geometry'], ensure_ascii=False)} as const;\n\n"
        f"export const REFERENCE_SAMPLES: RawReferenceSample[] = {json.dumps(samples, ensure_ascii=False, indent=1)};\n\n"
        f"export const MSA_SUMMARY: RawMsaSummary = {json.dumps(msa, ensure_ascii=False, indent=1)};\n\n"
        f"export const HUMAN_SUMMARY: RawHumanSummary | null = {json.dumps(human, ensure_ascii=False, indent=1)};\n\n"
        f"export const HUMAN_REPLIES: RawHumanReply[] = {json.dumps(human_replies, ensure_ascii=False, indent=1)};\n"
    )
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(body, encoding="utf-8")
    n_rep = sum(len(s["recorded"]) for s in samples)
    print(f"wrote {OUT.relative_to(REPO)}: {len(samples)} samples, {n_rep} recorded replies, "
          f"{len(human_replies)} human judgments")
    if human:
        c = human["combined"]
        print(f"human combined: own acc {c['own']['accuracy']:.1%} | eye+rules acc {c['rules']['accuracy']:.1%} | "
              f"fixed {c['fixedByRules']} broken {c['brokenByRules']} | repeatability own "
              f"{human['repeatability']['own']:.1%} rules {human['repeatability']['rules']:.1%} | "
              f"{human['secondsPerDie']} s/die")


if __name__ == "__main__":
    main()
