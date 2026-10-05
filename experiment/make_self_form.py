"""產生作者本人判定用的空白表單 results/self-judgment-form.xlsx。

- 第1輪／第2輪：40 張樣本以不同亂數順序排列，依檢驗規範 v1 填寫（兩輪建議隔天進行）。
- 重複性計算：以公式比對兩輪判定，算出本人重複性、對標準答案一致率、漏判率、誤判率。
- 標準答案（隱藏工作表）：兩輪都填完後再取消隱藏查看。

Usage: python experiment/make_self_form.py
"""
from __future__ import annotations

import json
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
HDR = Font(bold=True, color="FFFFFF")
HFILL = PatternFill("solid", fgColor="2F4F6F")
INPUT_FILL = PatternFill("solid", fgColor="FFF8DC")


def header(ws, values, widths):
    for j, (v, w) in enumerate(zip(values, widths), 1):
        c = ws.cell(row=1, column=j, value=v)
        c.font, c.fill = HDR, HFILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def round_sheet(wb, title, ids, seed):
    ws = wb.create_sheet(title)
    order = ids[:]
    random.Random(seed).shuffle(order)
    header(ws, ["序", "樣本", "缺陷代碼（無／CHP／CRK／SCR／CON，多個以 ; 分隔）", "區域（core／peripheral）",
                "量測 (µm)", "判定", "備註"], [6, 8, 34, 22, 14, 12, 30])
    n = len(order)
    for k, sid in enumerate(order, 2):
        ws.cell(row=k, column=1, value=k - 1)
        ws.cell(row=k, column=2, value=sid)
        for j in range(3, 8):
            ws.cell(row=k, column=j).fill = INPUT_FILL
    dv = DataValidation(type="list", formula1='"Pass,Warning,Fail"', allow_blank=True)
    dv.add(f"F2:F{n + 1}")
    ws.add_data_validation(dv)
    dz = DataValidation(type="list", formula1='"core,peripheral,core;peripheral"', allow_blank=True)
    dz.add(f"D2:D{n + 1}")
    ws.add_data_validation(dz)
    return ws


def main():
    ref = json.loads((HERE / "samples" / "reference.json").read_text(encoding="utf-8"))
    ids = [s["id"] for s in ref["samples"]]
    n = len(ids)
    wb = Workbook()
    ws = wb.active
    ws.title = "說明"
    for line in [
        "作者本人判定表（一致性分析：人工評估者）",
        "1. 依 docs/spec/inspection-spec-v1.md（檢驗規範 v1）判定 experiment/samples/ 中的 40 張樣本影像。",
        "2. 先填「第1輪」，隔天再填「第2輪」（兩輪順序已打亂且不同），第2輪時不要回看第1輪。",
        "3. 每張樣本一列：缺陷代碼、區域、量測 (µm，5 µm/px)、判定（下拉選 Pass／Warning／Fail）。無缺陷請填「無」並判 Pass。",
        "4. 「重複性計算」工作表會自動計算本人重複性、對標準答案一致率、漏判率、誤判率、人工複判率。",
        "5. 「標準答案」工作表已隱藏；兩輪都填完後再取消隱藏（右鍵工作表標籤 → 取消隱藏）。",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 110
    ws["A1"].font = Font(bold=True, size=13)

    round_sheet(wb, "第1輪", ids, seed=11)
    round_sheet(wb, "第2輪", ids, seed=22)

    rs = wb.create_sheet("標準答案")
    rs.append(["樣本", "標準答案判定", "標準答案缺陷"])
    for s in ref["samples"]:
        desc = "; ".join(
            f"{d['code']} {d['zone']} "
            + (f"{d.get('depthUm') or d.get('lengthUm') or d.get('diameterUm')}µm")
            + (" 經過pad" if d.get("crossesPad") else "") + (" 觸及seal ring" if d.get("touchesSealRing") else "")
            for d in s["defects"]) or "無"
        rs.append([s["id"], s["referenceVerdict"], desc])
    rs.column_dimensions["C"].width = 70
    rs.sheet_state = "hidden"

    cs = wb.create_sheet("重複性計算")
    header(cs, ["樣本", "第1輪判定", "第2輪判定", "兩輪一致 (1/0)", "標準答案", "第1輪對 (1/0)", "第2輪對 (1/0)"],
           [8, 12, 12, 14, 12, 14, 14])
    last = n + 1
    for k, sid in enumerate(sorted(ids), 2):
        cs.cell(row=k, column=1, value=sid)
        cs.cell(row=k, column=2, value=f'=IFERROR(IF(INDEX(第1輪!$F$2:$F${last},MATCH(A{k},第1輪!$B$2:$B${last},0))="","",INDEX(第1輪!$F$2:$F${last},MATCH(A{k},第1輪!$B$2:$B${last},0))),"")')
        cs.cell(row=k, column=3, value=f'=IFERROR(IF(INDEX(第2輪!$F$2:$F${last},MATCH(A{k},第2輪!$B$2:$B${last},0))="","",INDEX(第2輪!$F$2:$F${last},MATCH(A{k},第2輪!$B$2:$B${last},0))),"")')
        cs.cell(row=k, column=4, value=f'=IF(OR(B{k}="",C{k}=""),"",IF(B{k}=C{k},1,0))')
        cs.cell(row=k, column=5, value=f'=IFERROR(INDEX(標準答案!$B$2:$B${last},MATCH(A{k},標準答案!$A$2:$A${last},0)),"")')
        cs.cell(row=k, column=6, value=f'=IF(B{k}="","",IF(B{k}=E{k},1,0))')
        cs.cell(row=k, column=7, value=f'=IF(C{k}="","",IF(C{k}=E{k},1,0))')
    r = last + 2
    B, C, D, E, F, G = (f"{col}2:{col}{last}" for col in "BCDEFG")
    summary = [
        ("兩輪皆已填的樣本數", f"=COUNT({D})", "0"),
        ("重複性（兩輪判定一致比例）", f'=IF(COUNT({D})=0,"",SUM({D})/COUNT({D}))', "0.0%"),
        ("對標準答案一致率（第1輪）", f'=IF(COUNT({F})=0,"",SUM({F})/COUNT({F}))', "0.0%"),
        ("對標準答案一致率（第2輪）", f'=IF(COUNT({G})=0,"",SUM({G})/COUNT({G}))', "0.0%"),
        ("漏判率（標準答案 Fail 被判 Pass，兩輪合計）",
         f'=IF(COUNTIFS({E},"Fail",{B},"?*")+COUNTIFS({E},"Fail",{C},"?*")=0,"",(COUNTIFS({E},"Fail",{B},"Pass")+COUNTIFS({E},"Fail",{C},"Pass"))/(COUNTIFS({E},"Fail",{B},"?*")+COUNTIFS({E},"Fail",{C},"?*")))', "0.0%"),
        ("誤判率（標準答案 Pass 被判 Fail，兩輪合計）",
         f'=IF(COUNTIFS({E},"Pass",{B},"?*")+COUNTIFS({E},"Pass",{C},"?*")=0,"",(COUNTIFS({E},"Pass",{B},"Fail")+COUNTIFS({E},"Pass",{C},"Fail"))/(COUNTIFS({E},"Pass",{B},"?*")+COUNTIFS({E},"Pass",{C},"?*")))', "0.0%"),
        ("人工複判率（判 Warning 比例，兩輪合計）",
         f'=IF(COUNTIF({B},"?*")+COUNTIF({C},"?*")=0,"",(COUNTIF({B},"Warning")+COUNTIF({C},"Warning"))/(COUNTIF({B},"?*")+COUNTIF({C},"?*")))', "0.0%"),
    ]
    for k, (label, formula, fmt) in enumerate(summary):
        cs.cell(row=r + k, column=1, value=label).font = Font(bold=True)
        c = cs.cell(row=r + k, column=4, value=formula)
        c.number_format = fmt
    cs.column_dimensions["A"].width = 44
    out = HERE / "results" / "self-judgment-form.xlsx"
    out.parent.mkdir(exist_ok=True)
    wb.save(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
